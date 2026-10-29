import io
import json
import os
import random
import re
import secrets
import socket
import sqlite3
import time
import uuid
from datetime import date, timedelta
from difflib import SequenceMatcher
from pathlib import Path
from urllib.parse import quote

from flask import Flask, g, jsonify, redirect, render_template, request, session
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename
from pypdf import PdfReader
from docx import Document

BASE = Path(__file__).resolve().parent


def now():
    return time.time()


def create_app(config=None):
    app = Flask(__name__)
    app.config.update(SECRET_KEY=os.environ.get('STUDYSPACE_SECRET_KEY', 'studyspace-local-development-change-before-deployment'),
                      DATABASE=str(BASE / 'studyspace.db'), UPLOAD_FOLDER=str(BASE / 'uploads'),
                      MAX_CONTENT_LENGTH=10 * 1024 * 1024 + 65536,
                      PERMANENT_SESSION_LIFETIME=timedelta(hours=int(os.environ.get('SESSION_HOURS', '12'))),
                      SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax',
                      SESSION_COOKIE_NAME='studyspace_workspace_session')
    app.config['SEND_FILE_MAX_AGE_DEFAULT'] = 0
    if config:
        app.config.update(config)
    Path(app.config['UPLOAD_FOLDER']).mkdir(parents=True, exist_ok=True)

    def db():
        if 'db' not in g:
            g.db = sqlite3.connect(app.config['DATABASE'])
            g.db.row_factory = sqlite3.Row
            g.db.execute('PRAGMA foreign_keys=ON')
        return g.db

    @app.teardown_appcontext
    def close_db(error):
        if 'db' in g:
            g.db.close()

    with app.app_context():
        db().executescript('''
        CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY,name TEXT NOT NULL,email TEXT UNIQUE NOT NULL,password TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS documents(id INTEGER PRIMARY KEY,user_id INTEGER NOT NULL REFERENCES users(id),title TEXT NOT NULL,kind TEXT NOT NULL,filename TEXT NOT NULL,text TEXT NOT NULL,created TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS sections(id INTEGER PRIMARY KEY,document_id INTEGER REFERENCES documents(id) ON DELETE CASCADE,position INTEGER,text TEXT,completed INTEGER DEFAULT 0);
        CREATE TABLE IF NOT EXISTS questions(id INTEGER PRIMARY KEY,document_id INTEGER REFERENCES documents(id) ON DELETE CASCADE,prompt TEXT,options TEXT,answer TEXT,explanation TEXT,topic TEXT);
        CREATE TABLE IF NOT EXISTS attempts(id INTEGER PRIMARY KEY,user_id INTEGER REFERENCES users(id),document_id INTEGER REFERENCES documents(id) ON DELETE CASCADE,total INTEGER,correct INTEGER DEFAULT 0,submitted INTEGER DEFAULT 0,created TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS answers(id INTEGER PRIMARY KEY,attempt_id INTEGER REFERENCES attempts(id) ON DELETE CASCADE,question_id INTEGER REFERENCES questions(id) ON DELETE CASCADE,position INTEGER,options TEXT,selected TEXT,confidence TEXT,correct INTEGER);
        CREATE TABLE IF NOT EXISTS flashcards(id INTEGER PRIMARY KEY,user_id INTEGER REFERENCES users(id),question TEXT,answer TEXT,status TEXT DEFAULT 'new',reviews INTEGER DEFAULT 0);
        CREATE TABLE IF NOT EXISTS topics(id INTEGER PRIMARY KEY,user_id INTEGER REFERENCES users(id),title TEXT,due TEXT,completed INTEGER DEFAULT 0);
        CREATE TABLE IF NOT EXISTS study_sessions(id INTEGER PRIMARY KEY,user_id INTEGER REFERENCES users(id),token TEXT UNIQUE,started REAL,seconds INTEGER DEFAULT 0,completed INTEGER DEFAULT 0);
        CREATE TABLE IF NOT EXISTS activity(id INTEGER PRIMARY KEY,user_id INTEGER REFERENCES users(id),message TEXT,created TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS card_reviews(id INTEGER PRIMARY KEY,card_id INTEGER REFERENCES flashcards(id) ON DELETE CASCADE,user_id INTEGER REFERENCES users(id),status TEXT NOT NULL,created TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS mind_maps(document_id INTEGER PRIMARY KEY REFERENCES documents(id) ON DELETE CASCADE,data TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS reading_progress(document_id INTEGER PRIMARY KEY REFERENCES documents(id) ON DELETE CASCADE,user_id INTEGER REFERENCES users(id),section_id INTEGER REFERENCES sections(id) ON DELETE CASCADE,updated TEXT DEFAULT CURRENT_TIMESTAMP);
        ''')
        # Additive migrations keep accounts and progress from earlier versions intact.
        migrations = {
            'sections': {'title': "TEXT DEFAULT ''", 'blocks': "TEXT DEFAULT '[]'", 'format_version': 'INTEGER DEFAULT 0'},
            'flashcards': {'document_id': 'INTEGER REFERENCES documents(id) ON DELETE CASCADE', 'source': "TEXT DEFAULT ''", 'generated': 'INTEGER DEFAULT 0'},
            'topics': {'document_id': 'INTEGER REFERENCES documents(id) ON DELETE SET NULL'},
            'study_sessions': {'document_id': 'INTEGER REFERENCES documents(id) ON DELETE SET NULL', 'mode': "TEXT DEFAULT 'legacy'", 'duration': 'INTEGER DEFAULT 1500', 'elapsed': 'REAL DEFAULT 0', 'resumed': 'REAL', 'finished': 'TEXT'},
        }
        for table, columns in migrations.items():
            existing = {r['name'] for r in db().execute(f'PRAGMA table_info({table})')}
            for column, declaration in columns.items():
                if column not in existing:
                    db().execute(f'ALTER TABLE {table} ADD COLUMN {column} {declaration}')
        db().executescript('''
        CREATE INDEX IF NOT EXISTS documents_user ON documents(user_id);
        CREATE INDEX IF NOT EXISTS cards_user_doc ON flashcards(user_id,document_id);
        CREATE INDEX IF NOT EXISTS attempts_user_doc ON attempts(user_id,document_id);
        CREATE INDEX IF NOT EXISTS reviews_card ON card_reviews(card_id);
        CREATE INDEX IF NOT EXISTS topics_user_due ON topics(user_id,due);
        CREATE INDEX IF NOT EXISTS sessions_user ON study_sessions(user_id);
        CREATE INDEX IF NOT EXISTS sections_document_version ON sections(document_id,format_version,position);
        ''')
        for document in db().execute('SELECT * FROM documents').fetchall():
            ensure_reading_sections(db(), document)
        db().commit()

    def activity(message):
        db().execute('INSERT INTO activity(user_id,message) VALUES(?,?)', (session['uid'], message))

    def owned(table, item):
        if isinstance(item, bool) or not isinstance(item, (str, int)) or not str(item).isdigit():
            raise ValueError('Choose a valid record.')
        row = db().execute(f'SELECT * FROM {table} WHERE id=? AND user_id=?', (item, session['uid'])).fetchone()
        if not row:
            raise ValueError('That record was not found.')
        return row

    def field(name, maximum=500, minimum=1):
        value = (request.get_json(silent=True) or {}).get(name, '')
        if not isinstance(value, str) or not minimum <= len(value.strip()) <= maximum:
            raise ValueError(f'Enter a valid {name} ({minimum}–{maximum} characters).')
        return value.strip()

    @app.before_request
    def protect():
        if request.path.startswith('/api/'):
            if request.is_json and not isinstance(request.get_json(silent=True), dict):
                return jsonify(error='Send a valid object with your request.'), 400
            if request.path not in ['/api/session', '/api/login', '/api/register'] and 'uid' not in session:
                return jsonify(error='Please sign in to continue.'), 401
            if request.method != 'GET' and not secrets.compare_digest(request.headers.get('X-CSRF-Token', ''), session.get('csrf', secrets.token_hex())):
                return jsonify(error='Your session expired. Refresh the page and try again.'), 403

    @app.errorhandler(ValueError)
    def bad_input(error):
        return jsonify(error=str(error)), 400

    @app.errorhandler(413)
    def too_large(error):
        return jsonify(error='The maximum file size is 10 MB.'), 413

    @app.errorhandler(500)
    def server_error(error):
        return jsonify(error='Something went wrong. Please try again.'), 500

    @app.context_processor
    def asset_versions():
        # Changing CSS/JS changes its URL, so an embedded browser cannot retain old assets.
        return {'asset_version': max((BASE / 'static' / name).stat().st_mtime_ns
                                     for name in ('style.css', 'script.js'))}

    @app.after_request
    def fresh_pages(response):
        if request.path.startswith('/api/') or response.mimetype == 'text/html':
            response.headers['Cache-Control'] = 'no-store'
        response.headers['X-StudySpace-Edition'] = 'burgundy-workspace'
        return response

    @app.get('/')
    @app.get('/dashboard')
    @app.get('/materials')
    @app.get('/study')
    @app.get('/quizzes')
    @app.get('/forgetless')
    @app.get('/revision')
    @app.get('/timer')
    @app.get('/settings')
    def index():
        if 'uid' not in session:
            return redirect('/login?next=' + quote(request.full_path.rstrip('?'), safe='/'))
        return render_template('index.html')

    @app.get('/flashcards')
    @app.get('/flashcards/')
    @app.get('/flashcards/<path:old_path>')
    def retired_flashcards_page(old_path=None):
        # Keep saved records and API compatibility; the removed UI cannot be reopened.
        return redirect('/dashboard')

    @app.get('/login')
    def login_page():
        return redirect('/') if 'uid' in session else render_template('login.html')

    @app.get('/register')
    def register_page():
        return redirect('/') if 'uid' in session else render_template('register.html')

    @app.get('/api/session')
    def current_session():
        session.setdefault('csrf', secrets.token_hex(32))
        user = db().execute('SELECT id,name,email FROM users WHERE id=?', (session.get('uid'),)).fetchone()
        return jsonify(user=dict(user) if user else None, csrf=session['csrf'])

    @app.post('/api/register')
    def register():
        name, email, password = field('name', 80), field('email', 254).lower(), field('password', 128, 8)
        if not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', email):
            raise ValueError('Enter a valid email address.')
        try:
            with db():
                db().execute('INSERT INTO users(name,email,password) VALUES(?,?,?)', (name, email, generate_password_hash(password)))
        except sqlite3.IntegrityError:
            raise ValueError('This email is already registered. Please sign in.')
        return jsonify(ok=True)

    @app.post('/api/login')
    def login():
        email, password = field('email', 254).lower(), field('password', 128)
        user = db().execute('SELECT * FROM users WHERE email=?', (email,)).fetchone()
        if not user or not check_password_hash(user['password'], password):
            raise ValueError('Email or password is incorrect.')
        session.clear()
        session.update(uid=user['id'], csrf=secrets.token_hex(32))
        session.permanent = True
        return jsonify(ok=True)

    @app.post('/api/logout')
    def logout():
        session.clear()
        return jsonify(ok=True)

    @app.get('/api/data')
    def data():
        uid = session['uid']
        rows = lambda sql, args=(uid,): [dict(r) for r in db().execute(sql, args).fetchall()]
        docs = rows('SELECT d.id,d.title,d.kind,d.created,COUNT(s.id) sections,COALESCE(SUM(s.completed),0) completed,(SELECT p.updated FROM reading_progress p WHERE p.document_id=d.id) last_read,(SELECT r.position FROM reading_progress p JOIN sections r ON r.id=p.section_id WHERE p.document_id=d.id) reading_position FROM documents d LEFT JOIN sections s ON s.document_id=d.id AND s.format_version=1 WHERE d.user_id=? GROUP BY d.id ORDER BY d.id DESC')
        attempts = rows('SELECT a.*,d.title FROM attempts a JOIN documents d ON d.id=a.document_id WHERE a.user_id=? AND submitted=1 ORDER BY a.id DESC')
        gaps = rows("SELECT q.prompt,q.topic,q.explanation,d.title,d.id document_id FROM answers r JOIN attempts a ON a.id=r.attempt_id JOIN questions q ON q.id=r.question_id JOIN documents d ON d.id=a.document_id WHERE a.user_id=? AND a.submitted=1 AND r.confidence='High' AND r.correct=0 ORDER BY a.id DESC")
        insights = []
        for doc in docs:
            recent = next((a for a in attempts if a['document_id'] == doc['id']), None)
            has_gap = any(gap['document_id'] == doc['id'] for gap in gaps)
            if recent:
                score = 100 * recent['correct'] / recent['total']
                comparable = [a for a in attempts if a['document_id'] == doc['id'] and a['total'] == recent['total']]
                change = round(score - 100 * comparable[1]['correct'] / comparable[1]['total'], 1) if len(comparable) > 1 else None
                insights.append(dict(document_id=doc['id'], title=doc['title'], score=round(score, 1), change=change, total=recent['total'], status='Learning Gap' if has_gap else 'Needs Practice' if score < 60 else 'Developing' if score < 80 else 'Strong'))
            else:
                insights.append(dict(document_id=doc['id'], title=doc['title'], score=None, change=None, status='Not enough data yet'))
        sessions = rows("SELECT s.*,d.title document_title FROM study_sessions s LEFT JOIN documents d ON d.id=s.document_id WHERE s.user_id=? AND s.completed=1 ORDER BY s.id DESC LIMIT 100")
        week = []
        for offset in range(6, -1, -1):
            day = date.today() - timedelta(days=offset)
            seconds = db().execute("SELECT COALESCE(SUM(seconds),0) FROM study_sessions WHERE user_id=? AND completed=1 AND mode='focus' AND date(finished,'localtime')=?", (uid, day.isoformat())).fetchone()[0]
            week.append(dict(day=day.isoformat(), seconds=seconds))
        return jsonify(documents=docs, attempts=attempts, gaps=gaps, insights=insights,
                       cards=rows('SELECT * FROM flashcards WHERE user_id=? ORDER BY id DESC'),
                       topics=rows('SELECT * FROM topics WHERE user_id=? ORDER BY completed,due'),
                       activity=rows('SELECT message,created FROM activity WHERE user_id=? ORDER BY id DESC LIMIT 12'),
                       sessions=sessions, week=week,
                       practice=rows("SELECT c.id,c.question,c.document_id,COUNT(r.id) practice_count FROM flashcards c JOIN card_reviews r ON r.card_id=c.id WHERE c.user_id=? AND r.status='revise' GROUP BY c.id ORDER BY practice_count DESC LIMIT 10"),
                       study_seconds=db().execute("SELECT COALESCE(SUM(seconds),0) FROM study_sessions WHERE user_id=? AND completed=1 AND mode='focus'", (uid,)).fetchone()[0], today=date.today().isoformat())

    @app.post('/api/documents')
    def upload():
        file = request.files.get('file')
        if not file or not file.filename:
            raise ValueError('Choose a file first.')
        name = secure_filename(file.filename)
        kind = Path(name).suffix.lower().lstrip('.')
        if kind not in ('pdf', 'docx', 'txt'):
            raise ValueError('Please upload a PDF, DOCX, or TXT file.')
        raw = file.read(10 * 1024 * 1024 + 1)
        if len(raw) > 10 * 1024 * 1024:
            raise ValueError('The maximum file size is 10 MB.')
        try:
            if kind == 'pdf':
                text = '\n'.join(page.extract_text() or '' for page in PdfReader(io.BytesIO(raw)).pages)
            elif kind == 'docx':
                document = Document(io.BytesIO(raw))
                text = '\n'.join([p.text for p in document.paragraphs] + [' '.join(c.text for c in row.cells) for table in document.tables for row in table.rows])
            else:
                try:
                    text = raw.decode('utf-8-sig')
                except UnicodeDecodeError:
                    text = raw.decode('cp1252')
                if '\x00' in text:
                    raise ValueError()
        except Exception:
            raise ValueError('This file could not be read. Use an unencrypted PDF, a valid DOCX, or a text file.')
        text = re.sub(r'[ \t]+', ' ', text).strip()
        if len(re.findall(r'\w+', text)) < 8:
            raise ValueError('There is too little extractable text. Scanned PDFs need OCR before uploading; try a text-based file.')
        folder = Path(app.config['UPLOAD_FOLDER'])
        folder.mkdir(parents=True, exist_ok=True)
        filename = uuid.uuid4().hex + '.' + kind
        path = folder / filename
        path.write_bytes(raw)
        try:
            with db():
                docid = db().execute('INSERT INTO documents(user_id,title,kind,filename,text) VALUES(?,?,?,?,?)', (session['uid'], Path(name).stem[:150] or 'Study material', kind, filename, text)).lastrowid
                ensure_reading_sections(db(), db().execute('SELECT * FROM documents WHERE id=?', (docid,)).fetchone())
                generate_questions(db(), docid, text)
                activity('Uploaded ' + (Path(name).stem[:150] or 'study material'))
        except Exception:
            path.unlink(missing_ok=True)
            raise
        return jsonify(ok=True, id=docid)

    @app.route('/api/documents/<int:item>', methods=['GET', 'DELETE'])
    def document(item):
        doc = owned('documents', item)
        if request.method == 'GET':
            sections = [dict(r) for r in db().execute('SELECT * FROM sections WHERE document_id=? AND format_version=1 ORDER BY position', (item,))]
            for reading_section in sections:
                reading_section['blocks'] = json.loads(reading_section['blocks'])
            progress = db().execute('SELECT section_id FROM reading_progress WHERE document_id=? AND user_id=?', (item, session['uid'])).fetchone()
            resume = next((s['position'] for s in sections if progress and s['id'] == progress['section_id']), None)
            if resume is None:
                resume = next((s['position'] for s in sections if not s['completed']), 0)
            return jsonify(id=item, title=doc['title'], text=doc['text'], sections=sections, resume_position=resume,
                           concepts=extract_concepts(doc['text']),
                           questions=db().execute('SELECT COUNT(*) FROM questions WHERE document_id=?', (item,)).fetchone()[0], cards=db().execute('SELECT COUNT(*) FROM flashcards WHERE document_id=?', (item,)).fetchone()[0])
        with db():
            db().execute('DELETE FROM documents WHERE id=?', (item,))
            activity('Deleted ' + doc['title'])
        (Path(app.config['UPLOAD_FOLDER']) / Path(doc['filename']).name).unlink(missing_ok=True)
        return jsonify(ok=True)

    @app.post('/api/documents/<int:item>/position')
    def save_reading_position(item):
        owned('documents', item)
        section_id = (request.get_json(silent=True) or {}).get('section_id')
        if type(section_id) is not int:
            raise ValueError('Choose a valid reading section.')
        section = db().execute('SELECT id FROM sections WHERE id=? AND document_id=? AND format_version=1', (section_id, item)).fetchone()
        if not section:
            raise ValueError('This section does not belong to the selected material.')
        with db():
            db().execute('INSERT INTO reading_progress(document_id,user_id,section_id) VALUES(?,?,?) ON CONFLICT(document_id) DO UPDATE SET section_id=excluded.section_id,updated=CURRENT_TIMESTAMP', (item, session['uid'], section_id))
        return jsonify(ok=True)

    @app.post('/api/documents/<int:item>/flashcards')
    def generate_cards(item):
        doc = owned('documents', item)
        payload = request.get_json(silent=True) or {}
        target = payload.get('target', 10)
        if target not in (5, 10, 15, 'all'):
            raise ValueError('Choose 5, 10, 15, or all suitable cards.')
        pairs = extract_concepts(doc['text'])
        if not pairs:
            raise ValueError('No reliable definition or explanation patterns were found. Try clearer notes, or add a card manually.')
        existing = db().execute('SELECT COUNT(*) FROM flashcards WHERE document_id=? AND generated=1', (item,)).fetchone()[0]
        if existing and not payload.get('replace', False):
            return jsonify(count=existing, preserved=True, message='Existing generated cards preserved. Use Regenerate to replace generated cards and their reviews; manual cards stay.')
        selected = pairs if target == 'all' else pairs[:target]
        with db():
            if payload.get('replace', False):
                db().execute('DELETE FROM flashcards WHERE document_id=? AND generated=1', (item,))
            for concept in selected:
                db().execute('INSERT INTO flashcards(user_id,document_id,question,answer,source,generated) VALUES(?,?,?,?,?,1)', (session['uid'], item, concept['question'], concept['answer'], concept['source']))
            activity(f'Generated {len(selected)} flashcards from {doc["title"]}')
        return jsonify(count=len(selected), preserved=False, message=f'{len(selected)} source-based cards generated. The source supports {len(pairs)} suitable cards.')

    @app.route('/api/documents/<int:item>/mindmap', methods=['GET', 'POST'])
    def mindmap(item):
        doc = owned('documents', item)
        row = db().execute('SELECT data FROM mind_maps WHERE document_id=?', (item,)).fetchone()
        if row and (request.method == 'GET' or not (request.get_json(silent=True) or {}).get('replace')):
            return jsonify(json.loads(row['data']))
        if request.method == 'GET':
            return jsonify(generated=False, nodes=[], title=doc['title'])
        result = build_mindmap(doc['title'], doc['text'])
        with db():
            db().execute('INSERT INTO mind_maps(document_id,data) VALUES(?,?) ON CONFLICT(document_id) DO UPDATE SET data=excluded.data', (item, json.dumps(result)))
            activity('Generated a mind map from ' + doc['title'])
        return jsonify(result)

    @app.post('/api/sections/<int:item>')
    def complete_section(item):
        row = db().execute('SELECT s.* FROM sections s JOIN documents d ON d.id=s.document_id WHERE s.id=? AND d.user_id=?', (item, session['uid'])).fetchone()
        if not row:
            raise ValueError('Section not found.')
        with db():
            db().execute('UPDATE sections SET completed=1 WHERE id=?', (item,))
            if not row['completed']:
                activity('Completed a learning section')
        return jsonify(ok=True)

    @app.post('/api/quizzes')
    def start_quiz():
        payload = request.get_json() or {}
        doc = owned('documents', payload.get('document_id'))
        length = payload.get('length')
        if length not in (5, 10, 15):
            raise ValueError('Choose 5, 10, or 15 questions.')
        questions = db().execute('SELECT * FROM questions WHERE document_id=?', (doc['id'],)).fetchall()
        random.shuffle(questions)
        questions = questions[:length]
        if not questions:
            raise ValueError('No reliable questions could be generated. Try material with at least four clear definitions such as “Photosynthesis is …”.')
        with db():
            attempt = db().execute('INSERT INTO attempts(user_id,document_id,total) VALUES(?,?,?)', (session['uid'], doc['id'], len(questions))).lastrowid
            result = []
            for pos, q in enumerate(questions):
                options = json.loads(q['options'])
                random.shuffle(options)
                aid = db().execute('INSERT INTO answers(attempt_id,question_id,position,options) VALUES(?,?,?,?)', (attempt, q['id'], pos, json.dumps(options))).lastrowid
                result.append(dict(id=aid, prompt=q['prompt'], options=options, selected=None, confidence=None))
        return jsonify(id=attempt, questions=result)

    @app.post('/api/quizzes/<int:item>/submit')
    def submit_quiz(item):
        attempt = owned('attempts', item)
        if attempt['submitted']:
            raise ValueError('This quiz has already been submitted.')
        supplied = (request.get_json() or {}).get('answers', [])
        if not isinstance(supplied, list):
            raise ValueError('Invalid answers.')
        supplied = {a.get('id'): a for a in supplied if isinstance(a, dict)}
        rows = db().execute('SELECT a.*,q.answer,q.explanation,q.prompt FROM answers a JOIN questions q ON q.id=a.question_id WHERE attempt_id=? ORDER BY position', (item,)).fetchall()
        results, score = [], 0
        with db():
            for row in rows:
                choice = supplied.get(row['id'], {})
                if choice.get('selected') not in json.loads(row['options']) or choice.get('confidence') not in ('Low', 'Medium', 'High'):
                    raise ValueError('Answer every question and choose your confidence before submitting.')
                correct = int(choice['selected'] == row['answer'])
                score += correct
                db().execute('UPDATE answers SET selected=?,confidence=?,correct=? WHERE id=?', (choice['selected'], choice['confidence'], correct, row['id']))
                results.append(dict(prompt=row['prompt'], answer=row['answer'], explanation=row['explanation'], selected=choice['selected'], confidence=choice['confidence'], correct=correct))
            db().execute('UPDATE attempts SET correct=?,submitted=1 WHERE id=?', (score, item))
            activity(f'Finished a quiz: {score}/{len(rows)} correct')
        return jsonify(correct=score, total=len(rows), results=results)

    @app.route('/api/cards', methods=['POST'])
    @app.route('/api/cards/<int:item>', methods=['PUT', 'DELETE'])
    def cards(item=None):
        if item:
            owned('flashcards', item)
        with db():
            if request.method == 'DELETE':
                db().execute('DELETE FROM flashcards WHERE id=?', (item,))
            elif item and 'status' in (request.get_json() or {}):
                status = request.json['status']
                if status not in ('known', 'revise'):
                    raise ValueError('Invalid review status.')
                db().execute('UPDATE flashcards SET status=?,reviews=reviews+1 WHERE id=?', (status, item))
                db().execute('INSERT INTO card_reviews(card_id,user_id,status) VALUES(?,?,?)', (item, session['uid'], status))
                activity('Reviewed a flashcard')
            else:
                question, answer = field('question', 1000), field('answer', 5000)
                document_id = (request.get_json() or {}).get('document_id')
                if document_id is not None:
                    owned('documents', document_id)
                if item:
                    db().execute('UPDATE flashcards SET question=?,answer=?,document_id=? WHERE id=?', (question, answer, document_id, item))
                else:
                    db().execute('INSERT INTO flashcards(user_id,question,answer,document_id) VALUES(?,?,?,?)', (session['uid'], question, answer, document_id))
                    activity('Created a flashcard')
        return jsonify(ok=True)

    @app.route('/api/topics', methods=['POST'])
    @app.route('/api/topics/<int:item>', methods=['PUT', 'DELETE'])
    def topics(item=None):
        if item:
            owned('topics', item)
        with db():
            if request.method == 'DELETE':
                db().execute('DELETE FROM topics WHERE id=?', (item,))
            elif item and 'completed' in (request.get_json() or {}):
                db().execute('UPDATE topics SET completed=? WHERE id=?', (int(bool(request.json['completed'])), item))
            else:
                title, due = field('title', 200), field('due', 10)
                document_id = (request.get_json() or {}).get('document_id')
                if document_id is not None:
                    owned('documents', document_id)
                try:
                    date.fromisoformat(due)
                except ValueError:
                    raise ValueError('Choose a valid revision date.')
                if item:
                    db().execute('UPDATE topics SET title=?,due=?,document_id=? WHERE id=?', (title, due, document_id, item))
                else:
                    db().execute('INSERT INTO topics(user_id,title,due,document_id) VALUES(?,?,?,?)', (session['uid'], title, due, document_id))
                    activity('Planned revision: ' + title)
        return jsonify(ok=True)

    @app.post('/api/timer/start')
    def timer_start():
        payload = request.get_json(silent=True) or {}
        mode = payload.get('mode', 'focus')
        if mode not in ('focus', 'short', 'long'):
            raise ValueError('Choose a valid timer mode.')
        document_id = payload.get('document_id')
        if document_id is not None:
            owned('documents', document_id)
        token = uuid.uuid4().hex
        duration = {'focus': 1500, 'short': 300, 'long': 900}[mode]
        with db():
            db().execute('INSERT INTO study_sessions(user_id,token,started,resumed,mode,duration,document_id) VALUES(?,?,?,?,?,?,?)', (session['uid'], token, now(), now(), mode, duration, document_id))
        return jsonify(token=token, duration=duration)

    @app.post('/api/timer/control')
    def timer_control():
        payload = request.get_json(silent=True) or {}
        row = db().execute('SELECT * FROM study_sessions WHERE token=? AND user_id=?', (payload.get('token'), session['uid'])).fetchone()
        if not row or row['completed']:
            raise ValueError('No active timer found.')
        action = payload.get('action')
        elapsed, resumed = row['elapsed'], row['resumed']
        if action == 'pause' and resumed is not None:
            elapsed += max(0, now() - resumed)
            resumed = None
        elif action == 'resume' and resumed is None:
            resumed = now()
        elif action not in ('pause', 'resume'):
            raise ValueError('Invalid timer action.')
        with db():
            db().execute('UPDATE study_sessions SET elapsed=?,resumed=? WHERE id=?', (elapsed, resumed, row['id']))
        return jsonify(elapsed=min(row['duration'], elapsed), running=resumed is not None)

    @app.post('/api/timer/finish')
    def timer_finish():
        payload = request.get_json() or {}
        row = db().execute('SELECT * FROM study_sessions WHERE token=? AND user_id=?', (payload.get('token'), session['uid'])).fetchone()
        if not row:
            raise ValueError('Start a study session first.')
        elapsed = row['elapsed'] + (max(0, now() - row['resumed']) if row['resumed'] is not None else 0)
        if not row['completed'] and elapsed < row['duration']:
            raise ValueError('Finish the full countdown before saving. Incomplete sessions do not count.')
        with db():
            if not row['completed']:
                db().execute('UPDATE study_sessions SET seconds=?,completed=1,finished=CURRENT_TIMESTAMP,resumed=NULL WHERE id=?', (row['duration'], row['id']))
                activity(f'Completed a {row["duration"] // 60}-minute {row["mode"]} session')
        return jsonify(ok=True)

    @app.post('/api/settings')
    def settings():
        name = field('name', 80)
        with db():
            db().execute('UPDATE users SET name=? WHERE id=?', (name, session['uid']))
        return jsonify(ok=True)

    return app


def reading_blocks(text):
    """Format source lines without summarising, reordering, or changing their wording."""
    lines = text.replace('\r\n', '\n').replace('\r', '\n').replace('\f', '\n\n').split('\n')
    blocks, paragraph = [], []

    def flush():
        if paragraph:
            source = '\n'.join(paragraph)
            blocks.append(dict(kind='paragraph', text=' '.join(line.strip() for line in paragraph), source=source))
            paragraph.clear()

    for index, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        bullet = re.match(r'^(?:[\u2022\u25cf\u25aa*\-])\s+(.+)$', stripped)
        numbered = re.match(r'^(\d{1,3})[.)]\s+(.+)$', stripped)
        explicit_heading = re.match(r'^(?:#{1,4}\s+|(?:chapter|unit|module|section|part)\s+\d+\b)', stripped, re.I)
        heading_candidate = (len(stripped) <= 85 and len(stripped.split()) <= 10
                             and not re.search(r'[.!?;,]$', stripped)
                             and not re.search(r'\b(is|are|was|were|means|refers|consists)\b', stripped, re.I))
        isolated_title = (heading_candidate and stripped[0].isupper()
                          and (index == 0 or not lines[index - 1].strip())
                          and (index + 1 == len(lines) or not lines[index + 1].strip()
                               or len(lines[index + 1].split()) > 10))
        is_heading = bool(explicit_heading or (heading_candidate and stripped.isupper()) or isolated_title)
        if is_heading and not bullet:
            flush()
            blocks.append(dict(kind='heading', text=re.sub(r'^#{1,4}\s+', '', stripped), source=line))
        elif bullet or numbered:
            flush()
            blocks.append(dict(kind='bullet' if bullet else 'number', text=bullet[1] if bullet else numbered[2],
                               source=line, number=int(numbered[1]) if numbered else None))
        elif line[:1].isspace() and blocks and blocks[-1]['kind'] in ('bullet', 'number') and not paragraph:
            blocks[-1]['text'] += ' ' + stripped
            blocks[-1]['source'] += '\n' + line
        else:
            paragraph.append(line)
    flush()
    # Very long paragraphs may be divided only between complete source sentences.
    result = []
    for block in blocks:
        if block['kind'] != 'paragraph' or len(block['text'].split()) <= 420:
            result.append(block)
            continue
        sentences = re.split(r'(?<=[.!?])\s+(?=[A-Z0-9\"\u201c])', block['text'])
        current = []
        for sentence in sentences:
            current.append(sentence)
            if len(' '.join(current).split()) >= 300:
                passage = ' '.join(current)
                result.append(dict(kind='paragraph', text=passage, source=passage))
                current = []
        if current:
            passage = ' '.join(current)
            result.append(dict(kind='paragraph', text=passage, source=passage))
    return result


def split_reading_sections(text):
    blocks = reading_blocks(text)
    if not blocks:
        return []
    # Keep each source heading attached to the following passage/list item.
    units, pending = [], []
    for block in blocks:
        pending.append(block)
        if block['kind'] != 'heading':
            units.append(pending)
            pending = []
    if pending:
        if units:
            units[-1].extend(pending)
        else:
            units.append(pending)
    weights = [sum(len(b['source'].split()) for b in unit) for unit in units]
    total = sum(weights)
    target = min(8, max(6, (total + 649) // 650)) if total >= 1200 else max(1, (total + 349) // 350)
    count = min(target, len(units))
    cumulative = [0]
    for weight in weights:
        cumulative.append(cumulative[-1] + weight)
    boundaries, start = [], 0
    for part in range(1, count):
        stop = min(range(start + 1, len(units) - (count - part) + 1),
                   key=lambda boundary: abs(cumulative[boundary] - total * part / count))
        boundaries.append((start, stop))
        start = stop
    boundaries.append((start, len(units)))
    result = []
    for start, stop in boundaries:
        section_blocks = [block for unit in units[start:stop] for block in unit]
        heading = next((b['text'] for b in section_blocks if b['kind'] == 'heading'), '')
        result.append(dict(title=heading, blocks=section_blocks,
                           text='\n\n'.join(b['source'] for b in section_blocks)))
    return result


def ensure_reading_sections(connection, document):
    """Retain legacy rows; carry completion forward only for fully covered source spans."""
    if connection.execute('SELECT 1 FROM sections WHERE document_id=? AND format_version=1', (document['id'],)).fetchone():
        return
    legacy = connection.execute('SELECT text,completed FROM sections WHERE document_id=? ORDER BY position,id', (document['id'],)).fetchall()
    previous_words = [word for row in legacy for word in row['text'].split()]
    compatible = previous_words == document['text'].split()
    completed_ranges, offset = [], 0
    for row in legacy:
        end = offset + len(row['text'].split())
        if row['completed']:
            completed_ranges.append((offset, end))
        offset = end
    offset = 0
    for position, section in enumerate(split_reading_sections(document['text'])):
        end = offset + len(section['text'].split())
        covered = sum(max(0, min(end, finish) - max(offset, begin)) for begin, finish in completed_ranges)
        completed = int(compatible and covered == end - offset and end > offset)
        connection.execute('INSERT INTO sections(document_id,position,text,completed,title,blocks,format_version) VALUES(?,?,?,?,?,?,1)',
                           (document['id'], position, section['text'], completed, section['title'], json.dumps(section['blocks'])))
        offset = end


def extract_concepts(text):
    """Conservative English patterns; every answer retains its source passage."""
    concepts, seen = [], set()
    term_pattern = r'([A-Za-z][A-Za-z0-9() /\-]{1,65}?)'
    patterns = [
        (r'The main function of ' + term_pattern + r' is (.+)', 'What is the main function of {}?'),
        (term_pattern + r' (?:is|are) defined as (.+)', 'What is {}?'),
        (term_pattern + r' refers to (.+)', 'What does {} refer to?'),
        (term_pattern + r' consists of (.+)', 'What does {} consist of?'),
        (term_pattern + r' (?:is|are|means) (.+)', 'What is {}?'),
        (term_pattern + r' (?:helps|enables) (.+)', 'What does {} help or enable?'),
    ]
    for passage in re.split(r'(?<=[.!?])\s+|\n+', text):
        passage = re.sub(r'\s+', ' ', passage).strip(' \u2022\t')
        for pattern, question in patterns:
            match = re.fullmatch(pattern, passage.rstrip('.!?'), flags=re.I)
            if not match:
                continue
            term, answer = (part.strip() for part in match.groups())
            key = re.sub(r'[^a-z0-9]', '', re.sub(r'^(?:the|an?)\s+', '', term.lower()))
            answer_key = re.sub(r'\W+', ' ', answer.lower()).strip()
            near_duplicate = any(SequenceMatcher(None, answer_key, re.sub(r'\W+', ' ', c['answer'].lower()).strip()).ratio() >= .97 for c in concepts)
            if (key in seen or len(term.split()) > 9 or len(answer.split()) < 5
                    or near_duplicate or len(answer) > 700 or term.lower() in ('it', 'this', 'that', 'there', 'these', 'they', 'which')):
                break
            seen.add(key)
            concepts.append(dict(term=term, question=question.format(term), answer=answer, source=passage))
            break
    return concepts[:100]


def build_mindmap(title, text):
    """Heading membership follows source order, never an inferred relationship."""
    concepts = extract_concepts(text)
    headings = []
    for line in text.splitlines():
        cleaned = line.strip().lstrip('#').strip()
        if (2 <= len(cleaned) <= 70 and len(cleaned.split()) <= 8
                and not re.search(r'[.!?]$', cleaned)
                and not re.search(r'\b(is|are|refers|consists|means)\b', cleaned, re.I)):
            headings.append((cleaned.rstrip(':'), text.find(line)))
    branches = []
    if headings:
        for label, offset in headings[:12]:
            following = next((pos for _, pos in headings if pos > offset), len(text))
            children = [dict(label=c['term'], source=c['source']) for c in concepts if offset <= text.find(c['source']) < following][:10]
            branches.append(dict(label=label, source=text[offset:following].strip()[:1800], children=children))
        for concept in concepts:
            if text.find(concept['source']) < headings[0][1]:
                branches.append(dict(label=concept['term'], source=concept['source'], children=[]))
    else:
        branches = [dict(label=c['term'], source=c['source'], children=[]) for c in concepts[:18]]
    limitation = 'Branches group concepts under the preceding source heading. Connections show source organization, not inferred relationships.' if headings else 'No reliable heading hierarchy was found. Concepts connect directly to the document; no deeper relationships are inferred.' if branches else 'No reliable concepts or headings were found. Try notes with clear headings and complete definitions.'
    return dict(generated=True, title=title, nodes=branches, limitation=limitation)


def generate_questions(connection, docid, text):
    concepts = extract_concepts(text)
    existing = {r[0] for r in connection.execute('SELECT prompt FROM questions WHERE document_id=?', (docid,))}
    for concept in concepts:
        alternatives = [c['term'] for c in concepts if c['term'] != concept['term'] and c['answer'].casefold() != concept['answer'].casefold()]
        prompt = 'Which source concept matches this explanation?\n' + concept['answer']
        if len(alternatives) < 3 or prompt in existing:
            continue
        options = [concept['term']] + alternatives[:3]
        connection.execute('INSERT INTO questions(document_id,prompt,options,answer,explanation,topic) VALUES(?,?,?,?,?,?)',
                           (docid, prompt, json.dumps(options), concept['term'], concept['source'], concept['term']))


app = create_app()

if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser(description='Run StudySpace from this project folder.')
    parser.add_argument('--port', type=int, default=5000)
    args = parser.parse_args()
    if not 1024 <= args.port <= 65535:
        parser.error('Choose a port from 1024 to 65535.')
    port = args.port
    while port <= 65535:
        with socket.socket() as probe:
            try:
                probe.bind(('127.0.0.1', port))
                break
            except OSError:
                port += 1
    if port > 65535:
        parser.error('No free local port was found.')
    print(f'Project: {BASE}', flush=True)
    if port != args.port:
        print(f'Port {args.port} is occupied by another app. Using port {port}.', flush=True)
    print(f'Open this StudySpace: http://127.0.0.1:{port}/dashboard', flush=True)
    app.run(host='127.0.0.1', port=port, debug=False)
